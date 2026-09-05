from docx import Document

class DocxService:
    def __init__(self, docx_path):
        self.docx_path = docx_path

    def extract_text(self):
        doc = Document(self.docx_path)
        text = ""

        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"

        return text

    def extract_tables(self):
        doc = Document(self.docx_path)
        tables = []

        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = [cell.text for cell in row.cells]
                table_data.append(row_data)
            tables.append(table_data)

        return tables

    